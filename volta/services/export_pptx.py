"""PowerPoint brief of the overview: key figures and the main charts.

Everything is drawn natively (text boxes, shapes and PowerPoint charts) so
the deck can be edited after export. Colours follow the dashboard theme.
"""

from __future__ import annotations

import io
from datetime import date
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from .kpis import fmt_compact, fmt_full

NAVY = RGBColor(0x0B, 0x25, 0x45)
BLUE = RGBColor(0x00, 0x60, 0xA8)
RED = RGBColor(0xC8, 0x39, 0x1F)
SKY = RGBColor(0x00, 0x98, 0xD8)
GOLD = RGBColor(0xE0, 0xA0, 0x00)
INK = RGBColor(0x0B, 0x0B, 0x0B)
INK2 = RGBColor(0x52, 0x51, 0x4E)
MUTED = RGBColor(0x89, 0x87, 0x81)
PLANE = RGBColor(0xF3, 0xF5, 0xF8)
HAIRLINE = RGBColor(0xE1, 0xE0, 0xD9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x1C, 0x7C, 0x3A)
SERIES = [BLUE, RED, SKY, GOLD]

FONT = "Calibri"
SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6)

LOGO_PATH = Path(__file__).resolve().parents[1] / "static" / "img" / "nedco.png"


# ------------------------------------------------------------------ helpers
def _text(slide, left, top, width, height, text, *, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT,
          anchor=MSO_ANCHOR.TOP, font=FONT, italic=False):
    box = slide.shapes.add_textbox(int(left), int(top), int(width), int(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    f = run.font
    f.name, f.size, f.bold, f.italic = font, Pt(size), bold, italic
    f.color.rgb = color
    return box


def _add_line(paragraph_box, text, *, size=14, bold=False, color=INK, space_before=0):
    """Append a paragraph to an existing text box."""
    tf = paragraph_box.text_frame
    p = tf.add_paragraph()
    p.space_before = Pt(space_before)
    run = p.add_run()
    run.text = text
    f = run.font
    f.name, f.size, f.bold = FONT, Pt(size), bold
    f.color.rgb = color
    return p


def _card(slide, left, top, width, height, fill=PLANE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, int(left), int(top), int(width), int(height))
    shape.adjustments[0] = 0.06
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _slide_title(slide, title, subtitle=None):
    _text(slide, MARGIN, Inches(0.45), SLIDE_W - 2 * MARGIN, Inches(0.6), title, size=28, bold=True, color=NAVY)
    if subtitle:
        _text(slide, MARGIN, Inches(1.1), SLIDE_W - 2 * MARGIN, Inches(0.4), subtitle, size=13, color=INK2)


def _footer(slide, text):
    _text(slide, MARGIN, SLIDE_H - Inches(0.55), SLIDE_W - 2 * MARGIN, Inches(0.3), text, size=9, color=MUTED)


def _fmt_date(d: Optional[date]) -> str:
    return d.strftime("%-d %b %Y") if d else "—"


def _delta_text(value: Optional[float]) -> tuple[str, RGBColor]:
    if value is None:
        return "no earlier period", MUTED
    if abs(value) <= 0.05:
        return "no change", INK2
    arrow = "▲" if value > 0 else "▼"
    return f"{arrow} {value:+.1f}%", (GREEN if value > 0 else RED)


def _style_axes(chart, *, value_format="#,##0", label_size=10):
    ca = chart.category_axis
    ca.tick_labels.font.size = Pt(label_size)
    ca.tick_labels.font.name = FONT
    ca.tick_labels.font.color.rgb = INK2
    ca.format.line.color.rgb = HAIRLINE
    ca.has_major_gridlines = False
    va = chart.value_axis
    va.tick_labels.font.size = Pt(label_size)
    va.tick_labels.font.name = FONT
    va.tick_labels.font.color.rgb = MUTED
    va.tick_labels.number_format = value_format
    va.tick_labels.number_format_is_linked = False
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = HAIRLINE
    va.major_gridlines.format.line.width = Pt(0.75)
    va.format.line.fill.background()


def _chart_title(chart, text):
    chart.has_title = True
    tf = chart.chart_title.text_frame
    tf.text = text
    p = tf.paragraphs[0]
    p.font.size, p.font.bold, p.font.name = Pt(14), True, FONT
    p.font.color.rgb = NAVY


# ------------------------------------------------------------------- slides
def _title_slide(prs, d: Dict[str, Any]):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # Right panel with the headline numbers.
    panel_left = Inches(7.6)
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, panel_left, 0, SLIDE_W - panel_left, SLIDE_H)
    panel.fill.solid()
    panel.fill.fore_color.rgb = NAVY
    panel.line.fill.background()
    panel.shadow.inherit = False

    if LOGO_PATH.exists():
        slide.shapes.add_picture(str(LOGO_PATH), MARGIN, Inches(0.6), height=Inches(1.1))
    _text(slide, MARGIN, Inches(2.0), Inches(6.5), Inches(1.5), "Customer sales brief", size=36, bold=True, color=NAVY)
    _text(slide, MARGIN, Inches(3.6), Inches(6.5), Inches(0.5), "Northern Electricity Distribution Company", size=16, color=INK2)

    box = _text(slide, MARGIN, Inches(4.4), Inches(6.5), Inches(2.0), "Period", size=11, bold=True, color=MUTED)
    _add_line(box, d["period_label"], size=16, color=INK)
    _add_line(box, "Filters", size=11, bold=True, color=MUTED, space_before=10)
    _add_line(box, d["filters_label"], size=14, color=INK)
    _footer(slide, f"Data through {d['data_through']} · generated {d['generated']} from the NEDCo customer dashboard")

    cur = d["kpis"]["current"]
    items = [
        ("Energy sold", fmt_compact(cur.get("energy")), "kWh"),
        ("Amount paid", fmt_compact(cur.get("paymoney")), "GH₵"),
        ("Active customers", fmt_compact(cur.get("customers")), ""),
    ]
    top = Inches(1.4)
    for label, value, unit in items:
        _text(slide, panel_left + Inches(0.7), top, Inches(4.6), Inches(0.35), label.upper(), size=11, bold=True,
              color=RGBColor(0xB9, 0xCF, 0xE3))
        _text(slide, panel_left + Inches(0.7), top + Inches(0.35), Inches(4.6), Inches(0.9),
              f"{value} {unit}".strip(), size=40, bold=True, color=WHITE)
        top += Inches(1.75)
    if not cur.get("purchases"):
        _text(slide, panel_left + Inches(0.7), Inches(6.2), Inches(4.6), Inches(0.6),
              "No transactions match these filters.", size=13, color=WHITE)


def _kpi_slide(prs, d: Dict[str, Any]):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    k = d["kpis"]
    cur, deltas, prev = k["current"], k["deltas"], k["previous"]
    compare = (
        f"Change versus {_fmt_date(prev['start'])} to {_fmt_date(prev['end'])}"
        + (" (earlier data only partly available)" if prev.get("partial") else "")
        if prev else "No earlier period to compare with"
    )
    _slide_title(slide, "Key figures", f"{d['period_label']} · {compare}")

    tiles = [
        ("Energy sold", fmt_compact(cur.get("energy")), "kWh", deltas.get("energy"),
         f"{cur.get('purchases', 0):,} prepaid purchases"),
        ("Amount paid", fmt_compact(cur.get("paymoney")), "GH₵", deltas.get("paymoney"), "revenue from prepaid purchases"),
        ("Active customers", fmt_compact(cur.get("customers")), "", deltas.get("customers"),
         f"{cur.get('districts') or 0} district{'s' if (cur.get('districts') or 0) != 1 else ''}"),
        ("Spend per customer", fmt_compact(cur.get("spend_per_customer_month")), "GH₵",
         deltas.get("spend_per_customer_month"), "per month with a purchase"),
        ("Average price", fmt_full(cur.get("price_per_kwh"), 2) if cur.get("price_per_kwh") is not None else "—",
         "GH₵/kWh", deltas.get("price_per_kwh"), "amount paid per kWh"),
        ("Residential customers",
         f"{cur['residential_share']:.0f}%" if cur.get("residential_share") is not None else "—", "",
         deltas.get("residential_share"), "share of active customers"),
    ]
    cols, rows = 3, 2
    gap = Inches(0.35)
    grid_top = Inches(1.75)
    tile_w = int((SLIDE_W - 2 * MARGIN - gap * (cols - 1)) / cols)
    tile_h = Inches(2.2)
    for i, (label, value, unit, delta, sub) in enumerate(tiles):
        r, c = divmod(i, cols)
        left = MARGIN + c * (tile_w + gap)
        top = grid_top + r * (tile_h + gap)
        _card(slide, left, top, tile_w, tile_h)
        pad = Inches(0.3)
        _text(slide, left + pad, top + pad, tile_w - 2 * pad, Inches(0.35), label, size=13, color=INK2)
        box = _text(slide, left + pad, top + Inches(0.65), tile_w - 2 * pad, Inches(0.8), value, size=34, bold=True, color=NAVY)
        if unit:
            run = box.text_frame.paragraphs[0].add_run()
            run.text = f"  {unit}"
            run.font.size, run.font.name, run.font.bold = Pt(14), FONT, False
            run.font.color.rgb = INK2
        dtext, dcolor = _delta_text(delta)
        line = _text(slide, left + pad, top + Inches(1.45), tile_w - 2 * pad, Inches(0.65), dtext, size=13,
                     bold=delta is not None, color=dcolor)
        _add_line(line, sub, size=11, color=MUTED, space_before=2)
    _footer(slide, "Amount paid is the revenue measure for prepaid customers. Spend per customer = amount paid ÷ "
                   "customer-months. Average price = amount paid ÷ energy sold.")


def _monthly_slide(prs, d: Dict[str, Any], key: str, title: str, unit: str, chart_type, color):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    labels: List[str] = d["monthly"]["labels"]
    values: List[float] = d["monthly"][key]
    total = sum(values)
    _slide_title(slide, title, f"Total per month, {d['period_label']} · {fmt_full(total)} {unit} in total")
    if not labels:
        _text(slide, MARGIN, Inches(3), SLIDE_W - 2 * MARGIN, Inches(1), "No transactions match these filters.",
              size=16, color=MUTED, align=PP_ALIGN.CENTER)
        return
    data = CategoryChartData()
    data.categories = labels
    data.add_series(title, values)
    frame = slide.shapes.add_chart(chart_type, MARGIN, Inches(1.7), SLIDE_W - 2 * MARGIN, Inches(4.9), data)
    chart = frame.chart
    chart.has_legend = False
    chart.font.name = FONT
    _style_axes(chart)
    if len(labels) > 24:
        chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    chart.has_title = False
    series = chart.plots[0].series[0]
    if chart_type == XL_CHART_TYPE.LINE_MARKERS or chart_type == XL_CHART_TYPE.LINE:
        series.format.line.color.rgb = color
        series.format.line.width = Pt(2.25)
        series.smooth = False
        series.marker.style = XL_MARKER_STYLE.NONE
    else:
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = color
        chart.plots[0].gap_width = 60
    _footer(slide, f"{unit} per calendar month for the current filters. Data through {d['data_through']}.")


def _district_slide(prs, d: Dict[str, Any]):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rows: List[Dict[str, Any]] = d["districts"]
    _slide_title(slide, "By district", f"Energy sold and amount paid, {d['period_label']}")
    if not rows:
        _text(slide, MARGIN, Inches(3), SLIDE_W - 2 * MARGIN, Inches(1), "No transactions match these filters.",
              size=16, color=MUTED, align=PP_ALIGN.CENTER)
        return
    top_rows = rows[:12]
    data = CategoryChartData()
    data.categories = [r["label"] for r in reversed(top_rows)]  # bars read top-down
    data.add_series("Energy (kWh)", [r["energy"] for r in reversed(top_rows)])
    chart_w = Inches(6.6)
    frame = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, MARGIN, Inches(1.7), chart_w, Inches(4.9), data)
    chart = frame.chart
    chart.has_legend = False
    chart.font.name = FONT
    _style_axes(chart)
    plot = chart.plots[0]
    plot.gap_width = 45
    series = plot.series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = BLUE
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = "#,##0", False
    dl.font.size, dl.font.name = Pt(9), FONT
    dl.font.color.rgb = INK2
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    _chart_title(chart, "Energy sold (kWh)")

    # Table on the right.
    table_left = MARGIN + chart_w + Inches(0.4)
    table_w = SLIDE_W - MARGIN - table_left
    n = len(top_rows) + 1
    shape = slide.shapes.add_table(n, 4, int(table_left), Inches(1.75), int(table_w), Inches(0.36) * n)
    table = shape.table
    widths = [0.31, 0.23, 0.23, 0.23]
    for i, w in enumerate(widths):
        table.columns[i].width = int(table_w * w)
    headers = ["District", "kWh", "GH₵", "Customers"]
    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        p = cell.text_frame.paragraphs[0]
        p.font.size, p.font.bold, p.font.name = Pt(10), True, FONT
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT
    for r, row in enumerate(top_rows, start=1):
        vals = [row["label"], fmt_full(row["energy"]), fmt_full(row["paymoney"]), f"{row['customers']:,}"]
        for c, v in enumerate(vals):
            cell = table.cell(r, c)
            cell.text = v
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r % 2 else PLANE
            p = cell.text_frame.paragraphs[0]
            p.font.size, p.font.name = Pt(10), FONT
            p.font.color.rgb = INK
            p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT
    note = f"Top {len(top_rows)} of {len(rows)} districts by energy sold." if len(rows) > len(top_rows) else \
        f"{len(rows)} district{'s' if len(rows) != 1 else ''} with transactions in the period."
    _footer(slide, note + " Customers = meters with at least one purchase.")


def _mix_slide(prs, d: Dict[str, Any]):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _slide_title(slide, "By account type", f"Share of active customers and of energy sold, {d['period_label']}")
    mix = d["mix"]
    if not mix["labels"]:
        _text(slide, MARGIN, Inches(3), SLIDE_W - 2 * MARGIN, Inches(1), "No transactions match these filters.",
              size=16, color=MUTED, align=PP_ALIGN.CENTER)
        return
    half = int((SLIDE_W - 2 * MARGIN - Inches(0.4)) / 2)
    for i, (series_key, title) in enumerate((("customers", "Share of active customers"), ("energy", "Share of energy sold"))):
        data = CategoryChartData()
        data.categories = mix["labels"]
        data.add_series(title, mix[series_key])
        left = MARGIN + i * (half + Inches(0.4))
        frame = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, int(left), Inches(1.7), half, Inches(4.9), data)
        chart = frame.chart
        chart.font.name = FONT
        _chart_title(chart, title)
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size, chart.legend.font.name = Pt(11), FONT
        chart.legend.font.color.rgb = INK2
        plot = chart.plots[0]
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.show_percentage, dl.show_value, dl.show_category_name = True, False, False
        dl.number_format, dl.number_format_is_linked = "0%", False
        dl.font.size, dl.font.bold, dl.font.name = Pt(12), True, FONT
        dl.font.color.rgb = WHITE
        for j, point in enumerate(plot.series[0].points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = SERIES[j % len(SERIES)]
            point.format.line.color.rgb = WHITE
    _footer(slide, "Customers = meters with at least one purchase in the period. Segments under 1% may be hidden by PowerPoint.")


def _about_slide(prs, d: Dict[str, Any]):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _slide_title(slide, "How to read these figures")
    col_w = int((SLIDE_W - 2 * MARGIN - Inches(0.5)) / 2)
    left_items = [
        ("Energy sold", "Sum of energy purchased (kWh) in the period."),
        ("Amount paid", "Sum of prepaid purchases (GH₵). The revenue measure for prepaid customers."),
        ("Active customers", "Meters with at least one purchase in the period."),
    ]
    right_items = [
        ("Spend per customer", "Amount paid divided by customer-months (months in which a customer bought at least once)."),
        ("Average price", "Amount paid divided by energy sold (GH₵ per kWh), a realised-tariff check."),
        ("Change versus the previous period", "Whole-month periods compare with the same months one period earlier; "
                                              "other ranges shift back by their own length."),
    ]
    for i, items in enumerate((left_items, right_items)):
        left = MARGIN + i * (col_w + Inches(0.5))
        box = _text(slide, left, Inches(1.6), col_w, Inches(4.5), "", size=12)
        first = True
        for label, desc in items:
            p = box.text_frame.paragraphs[0] if first else box.text_frame.add_paragraph()
            if not first:
                p.space_before = Pt(14)
            first = False
            run = p.add_run()
            run.text = label
            run.font.size, run.font.bold, run.font.name = Pt(15), True, FONT
            run.font.color.rgb = NAVY
            p2 = box.text_frame.add_paragraph()
            p2.space_before = Pt(2)
            run = p2.add_run()
            run.text = desc
            run.font.size, run.font.name = Pt(13), FONT
            run.font.color.rgb = INK2
    _footer(slide, f"Source: NEDCo customer dashboard, prepaid purchase records through {d['data_through']}. "
                   f"Filters: {d['filters_label']}. Generated {d['generated']}.")


# --------------------------------------------------------------------- build
def build_brief(d: Dict[str, Any]) -> bytes:
    """Build the deck from the gathered data and return the .pptx bytes."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    _title_slide(prs, d)
    _kpi_slide(prs, d)
    _monthly_slide(prs, d, "energy", "Energy sold per month", "kWh", XL_CHART_TYPE.LINE, BLUE)
    _monthly_slide(prs, d, "paymoney", "Amount paid per month", "GH₵", XL_CHART_TYPE.COLUMN_CLUSTERED, BLUE)
    _district_slide(prs, d)
    _mix_slide(prs, d)
    _about_slide(prs, d)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


__all__ = ["build_brief"]
